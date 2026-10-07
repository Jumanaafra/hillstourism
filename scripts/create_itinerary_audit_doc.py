from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path("docs/Hills_Tourism_Itinerary_Builder_Audit_and_Improvement.docx")

NAVY = "17365D"
BLUE = "2F75B5"
PALE = "EAF2F8"
LIGHT = "F5F7FA"
GRAY = "5B6573"
GREEN = "217346"
AMBER = "A65E00"
RED = "A61B1B"
BORDER = "D9D9D9"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_border(cell, color=BORDER, size="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)

def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.extend([c, u])
    t = OxmlElement("w:t")
    t.text = text
    run.extend([rPr, t])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

def add_para(doc, text="", style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

def add_numbered(doc, items):
    for idx, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.23)
        p.paragraph_format.first_line_indent = Inches(-0.23)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f"{idx}.  ").bold = True
        p.add_run(item)

def make_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[i])
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if r_idx % 2 == 1:
                set_cell_shading(cells[i], PALE)
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].text = str(value)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string("20242B")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size, before, after in (("Title", 25, 0, 10), ("Heading 1", 17, 14, 7), ("Heading 2", 12, 10, 4)):
    s = styles[name]
    s.font.name = "Aptos Display"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
styles["Title"].paragraph_format.space_after = Pt(6)
# Remove the border inherited from some Word built-in Title definitions.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
styles["List Bullet"].font.name = "Aptos"
styles["List Bullet"].font.size = Pt(10)
styles["List Bullet"].paragraph_format.space_after = Pt(3)
styles["List Number"].font.name = "Aptos"
styles["List Number"].font.size = Pt(10)
styles["List Number"].paragraph_format.space_after = Pt(4)

title = doc.add_paragraph(style="Title")
title.add_run("Hills Tourism Itinerary Builder Audit and Improvement Proposal")
title_ppr_direct = title._p.get_or_add_pPr()
title_border_direct = title_ppr_direct.find(qn("w:pBdr"))
if title_border_direct is not None:
    title_ppr_direct.remove(title_border_direct)
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(12)
r = sub.add_run("Product and data audit | 29 September 2026")
r.font.color.rgb = RGBColor.from_string(GRAY)
r.font.size = Pt(10)

intro = doc.add_paragraph()
intro.add_run("Main conclusion. ").bold = True
intro.add_run("The builder is already strong at producing a customized, client-ready itinerary PDF, but it behaves as a document editor rather than a trip-planning system. The highest-value improvement is a ")
r = intro.add_run("Trip Feasibility Check")
r.bold = True
intro.add_run(" that validates route time, attraction opening hours, closures, and schedule buffers before download. This single feature would reduce operational mistakes and make the itinerary more trustworthy for both staff and travelers.")

doc.add_heading("Scope and method", level=1)
add_para(doc, "The audit covers the admin Travel Itinerary Builder, the package itinerary model and validation, and the PDF export contract. Findings are based on a source-code review and the existing itinerary unit tests. The online comparison uses product documentation from Wanderlog, Google Maps Platform, and Kerala Tourism.")

doc.add_heading("Current feature audit", level=1)
make_table(doc, ["Area", "Current behavior", "Status"], [
    ("Starting flow", "Select an existing package or start a custom itinerary from scratch.", "Good"),
    ("Day editing", "Add, remove, reorder, and duplicate days; add timed activity slots and day images.", "Good"),
    ("Trip details", "Captures traveler counts, dates, destination, price, stay, vehicle, inclusions, and exclusions.", "Good"),
    ("PDF output", "Generates a branded, selectable-text A4 PDF in the browser.", "Good"),
    ("Package import", "Converts package activities into schedule slots at synthetic two-hour intervals beginning at 9:00 AM.", "Risk"),
    ("Validation", "PDF generation requires only a package name; dates, day titles, activity text, chronological order, and price consistency are not validated.", "Gap"),
    ("Feasibility", "No place identity, coordinates, route duration, opening hours, closure, ticket, or travel-buffer checks.", "Major gap"),
    ("Persistence", "Customized builder state is independent of the package record and is downloaded only; there is no draft save, version history, or enquiry attachment.", "Gap"),
    ("Captured but omitted", "Customer phone, email, pickup, and special requests are collected in the screen but are not included in the PDF data passed to the generator.", "Defect"),
    ("Dormant state", "Payment and theme values exist in component state/reset logic but are not connected to the export contract.", "Incomplete"),
], widths=[1.25, 4.45, 0.8])

doc.add_heading("What is working well", level=2)
add_bullets(doc, [
    "The workflow supports both reusable package templates and fully custom itineraries.",
    "Day controls are practical for operations staff: duplicate, reorder, delete, image upload, and multiple schedule slots.",
    "Stay and vehicle information can be selected from existing CMS records or overridden for a specific customer.",
    "The package API validates unique positive day numbers plus required titles and descriptions, and the focused itinerary unit suite passed 14 tests during this audit.",
    "The generated PDF is client-friendly and does not modify the original package record.",
])

doc.add_heading("Important risks", level=2)
add_numbered(doc, [
    "False precision. Imported package activities receive invented times even though the source package stores no start time or visit duration.",
    "Impossible sequencing. A user can place distant attractions back-to-back without any transfer time or buffer.",
    "Closed attraction risk. A scheduled time can fall outside operating hours or on a seasonal closure date.",
    "Inconsistent dates. Travel dates, each day date, weekday, duration, and nights are separate free-text or numeric fields and can disagree.",
    "Quote mismatch. Estimated total cost is price per person multiplied by adults plus children, with no child pricing, tax, room, vehicle, ticket, or optional-activity breakdown.",
    "Information loss. Contact, pickup, and special-request fields shown to staff do not reach the PDF export data object.",
])

doc.add_heading("Online comparison", level=1)
add_para(doc, "Current travel-planning products treat place data and travel time as part of the itinerary, not as separate manual research. The comparison below highlights the capability gap relevant to Hills Tourism.")
make_table(doc, ["Reference", "Relevant capability", "Meaning for Hills Tourism"], [
    ("Wanderlog", "Shows plans on a map with distance and time between places and offers single-day route optimization.", "Confirms that route-aware day planning is a standard user expectation."),
    ("Google Routes API", "Computes distance and duration across multiple origins and destinations through a route matrix.", "Provides the technical basis for travel-time checks and stop ordering."),
    ("Google Places API", "Provides place identity, coordinates, business status, and regular/current opening hours including special days.", "Supports reliable place selection and open/closed validation."),
    ("Kerala Tourism", "Advises confirming conditions before visiting Eravikulam National Park and publishes visiting hours plus seasonal closure guidance.", "Shows why destination-specific closure warnings matter for Munnar itineraries."),
], widths=[1.3, 2.75, 2.9])

p = doc.add_paragraph()
p.add_run("Sources: ").bold = True
add_hyperlink(p, "Wanderlog route optimization", "https://help.wanderlog.com/hc/en-us/articles/13545624787867-Optimize-route")
p.add_run("; ")
add_hyperlink(p, "Google Routes route matrix", "https://developers.google.com/maps/documentation/routes/compute_route_matrix")
p.add_run("; ")
add_hyperlink(p, "Google Places resource and opening hours", "https://developers.google.com/maps/documentation/places/web-service/reference/rest/v1/places")
p.add_run("; ")
add_hyperlink(p, "Kerala Tourism Eravikulam fact file", "https://www.keralatourism.org/munnar/reservations-eravikulam-national-park.php")
p.add_run(".")

doc.add_heading("One recommended improvement", level=1)
doc.add_heading("Trip Feasibility Check", level=2)
add_para(doc, "Add a Validate Day button above each day and a Validate Entire Itinerary action beside Download PDF. The checker should resolve each activity to a structured place, calculate travel time between stops, compare arrival and visit windows with operating hours or known closure rules, and flag conflicts before the PDF is produced.")

doc.add_heading("How it should work", level=2)
add_numbered(doc, [
    "Staff selects a recognized place for each attraction while retaining a free-text note for pickups, meals, or custom stops.",
    "The system calculates transfer distance and duration from the previous stop, including the hotel or pickup point as the first origin when available.",
    "Each activity receives a planned arrival time and visit duration. The next start time is derived from travel plus visit duration, not guessed from a fixed interval.",
    "The checker shows blocking errors for closed or unreachable stops and warnings for tight buffers, missing booking references, or unverified seasonal information.",
    "Staff can accept an optimized order, keep the original order with an explanation, or edit the day. The PDF records travel legs and a last-checked timestamp.",
])

doc.add_heading("Missing data to add", level=2)
make_table(doc, ["Level", "Field", "Purpose"], [
    ("Activity", "placeId, placeName, latitude, longitude", "Identify the real location and avoid ambiguous free text."),
    ("Activity", "plannedStart, visitDurationMinutes", "Calculate arrival, departure, and conflicts."),
    ("Activity", "openingHoursSnapshot, businessStatus, checkedAt", "Explain the validation result and when it was verified."),
    ("Activity", "bookingRequired, bookingReference, ticketStatus", "Prevent scheduling a limited-entry attraction without the required reservation."),
    ("Travel leg", "fromPlaceId, toPlaceId, distanceMeters, durationSeconds", "Show realistic transfer requirements between stops."),
    ("Travel leg", "bufferMinutes, routeMode", "Allow mountain-road delays and vehicle-specific planning."),
    ("Day", "date as ISO value, startLocation, endLocation", "Replace independent date and weekday text with consistent structured data."),
    ("Day", "validationStatus, warnings, validatedAt", "Make readiness visible and auditable."),
    ("Trip", "timezone, pricing breakdown, currency", "Avoid date/time ambiguity and misleading totals."),
], widths=[0.9, 2.55, 3.5])

doc.add_heading("User interface result", level=2)
add_bullets(doc, [
    "Green: feasible, open, and enough buffer.",
    "Amber: feasible but tight, opening hours unverified, or booking not confirmed.",
    "Red: arrival after closing, route duration exceeds the available gap, seasonal closure, or unresolved place.",
    "Each schedule row shows travel time from the previous stop and a concise reason for any warning.",
    "Download remains available only after red issues are resolved or explicitly overridden with a staff note.",
])

doc.add_page_break()
doc.add_heading("Suggested implementation sequence", level=1)
make_table(doc, ["Stage", "Deliverable", "Acceptance check"], [
    ("1", "Structured activity and day schema plus date/time validation", "Dates, weekdays, duration, nights, and schedule order cannot disagree."),
    ("2", "Place search and place ID storage", "Every attraction can be resolved to coordinates while custom notes remain supported."),
    ("3", "Server-side route and opening-hours validation", "API keys remain server-only; errors return safe, actionable messages."),
    ("4", "Day status UI and optional optimized ordering", "Staff can see, fix, accept, or override every warning."),
    ("5", "PDF travel legs, warnings, and verification timestamp", "The exported itinerary reflects the validated plan and its check time."),
    ("6", "Draft persistence and enquiry linkage", "A customized itinerary can be reopened, versioned, and attached to the customer enquiry."),
], widths=[0.6, 3.0, 3.35])

doc.add_heading("Minimum viable release", level=2)
add_para(doc, "The first release does not need automatic AI itinerary generation. A focused MVP should support structured place selection, route duration, opening-hours checks, visit duration, a configurable mountain-road buffer, and clear red or amber warnings. This solves the core reliability problem while keeping staff in control.")

doc.add_heading("Acceptance criteria", level=2)
add_bullets(doc, [
    "A day cannot be marked ready when an attraction is closed at the planned arrival time.",
    "Travel time and visit duration are included when calculating the next activity start.",
    "Dates use a structured value and day-of-week is derived automatically.",
    "Route and Places credentials are never exposed to the browser.",
    "The PDF includes pickup, special requests, route legs, validation timestamp, and any approved override note.",
    "If an external provider is unavailable, staff sees an unverified status rather than a false success.",
    "Unit tests cover time overlap, closed venue, missing place, stale verification, and provider failure; an end-to-end test covers build, validate, and download.",
])

doc.add_heading("Decision", level=1)
add_para(doc, "Prioritize Trip Feasibility Check ahead of cosmetic builder enhancements. The current interface already creates attractive itineraries; the bigger business risk is sending a customer an itinerary whose timing, route, or attraction availability has never been validated. Fix the existing export omissions for pickup and special requests in the same delivery because they are already collected and require no new research workflow.")

# Footer with compact document identity and page field.
footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Hills Tourism itinerary builder audit  |  ")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string(GRAY)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
p._p.append(fld)

# Keep table rows together where practical and normalize fonts.
for table in doc.tables:
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.6)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = "Hills Tourism Itinerary Builder Audit and Improvement Proposal"
doc.core_properties.subject = "Audit and proposal for a trip feasibility check"
doc.core_properties.author = "OpenAI Codex"
doc.save(OUT)
print(OUT.resolve())
