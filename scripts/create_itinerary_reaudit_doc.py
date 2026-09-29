from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = Path('docs/Hills_Tourism_Itinerary_Builder_Reaudit_2026-09-29.docx')
NAVY = '102D52'
PALE = 'EAF4FC'
GRAY = '526170'

def shade(cell, fill):
    prop = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd')
    el.set(qn('w:fill'), fill)
    prop.append(el)

def cell_border(cell):
    prop = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement('w:' + side)
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '4')
        edge.set(qn('w:color'), 'D9D9D9')
        borders.append(edge)
    prop.append(borders)

def cell_padding(cell):
    prop = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for side in ('top', 'start', 'bottom', 'end'):
        edge = OxmlElement('w:' + side)
        edge.set(qn('w:w'), '110')
        edge.set(qn('w:type'), 'dxa')
        margins.append(edge)
    prop.append(margins)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, header in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Inches(widths[i])
        shade(c, NAVY)
        cell_border(c)
        cell_padding(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    row_prop = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    row_prop.append(repeat)
    for index, data in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(data):
            cells[i].width = Inches(widths[i])
            cells[i].text = value
            if index % 2:
                shade(cells[i], PALE)
            cell_border(cells[i])
            cell_padding(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8.5)
        prop = t.rows[-1]._tr.get_or_add_trPr()
        prop.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph()

def paragraph(text, lead=None):
    p = doc.add_paragraph()
    if lead and text.startswith(lead):
        p.add_run(lead).bold = True
        p.add_run(text[len(lead):])
    else:
        p.add_run(text)
    return p

def bullets(items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75)
sec.right_margin = Inches(.75)
base = doc.styles['Normal']
base.font.name = 'Aptos'
base.font.size = Pt(10)
base.font.color.rgb = RGBColor(32, 38, 46)
base.paragraph_format.space_after = Pt(6)
base.paragraph_format.line_spacing = 1.1
for name, size in (('Title', 23), ('Heading 1', 16), ('Heading 2', 11)):
    s = doc.styles[name]
    s.font.name = 'Aptos Display'
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True
    ppr = s._element.get_or_add_pPr()
    border = ppr.find(qn('w:pBdr'))
    if border is not None:
        ppr.remove(border)
doc.styles['List Bullet'].paragraph_format.space_after = Pt(3)

doc.add_paragraph('Hills Tourism Itinerary Builder Reaudit and Fix Report', style='Title')
p = doc.add_paragraph('29 September 2026  |  Admin itinerary workflow and customer PDF')
p.paragraph_format.space_after = Pt(12)
for run in p.runs:
    run.font.color.rgb = RGBColor.from_string(GRAY)

paragraph('Result. The package editor now opens in a popup when staff selects a package. Staff can edit the per person amount and other trip details there, choose a Sky Blue or Midnight Mountain theme, and download a PDF that uses those edits. This report records the original failure, the code changes, verification, and remaining gaps.', 'Result. ')

doc.add_heading('What the repeat audit found', level=1)
table(['Finding', 'Evidence in the updated code', 'Impact before fix'], [
    ('Editor location', 'The selected package populated a long editor below the package cards.', 'Staff had to scroll to find and review editable fields.'),
    ('PDF data loss', 'Phone, email, pickup, special requests, payment fields, and theme were held in component state but omitted from ItineraryPdfData.', 'Edited details did not appear in the downloaded PDF.'),
    ('Theme controls', 'Color and background controls existed in the screen but were not passed to the PDF generator.', 'The selected appearance had no effect on export.'),
    ('Price risk', 'The total used digits stripped from the full price string. A package advertising two rates could combine them into one amount.', 'The estimated total could be incorrect.'),
    ('Stale values', 'The editor began with example dates and default stay or vehicle text; selection did not reset every prior field.', 'A new itinerary could carry unrelated or outdated information.'),
], [1.25, 3.05, 2.15])

doc.add_heading('Changes completed', level=1)
bullets([
    'Selecting a package immediately opens a scrollable popup with a fixed header and a Download Edited PDF action. Closing it returns to the package list; the selected package can be reopened.',
    'The editable price per person drives the on-screen total and the PDF price and total. Multi-rate package strings require staff to choose one amount before download.',
    'The PDF now receives customer phone, email, pickup, special requests, infant count, and payment details when provided.',
    'Sky Blue and Midnight Mountain presets use the logo palette. The selected preset produces a matching vector cover in the PDF; an optional image can be added to that cover.',
    'Package selection clears prior customer and payment values. Existing packages without a linked stay or vehicle no longer inherit unrelated defaults.',
    'PDF money text uses INR so the downloaded file renders consistently with the bundled PDF font.',
])

doc.add_page_break()
doc.add_heading('Verified behavior', level=1)
table(['Check', 'Result'], [
    ('Browser flow', 'Local preview: package selection opened the popup; edited price changed estimated total from 4,598 to 9,000 for two adults; dark theme became selected.'),
    ('PDF export', 'Focused PDF test confirmed edited title, day, price, pickup, special requests, and payment text in the exported file; a dark themed cover rendered in visual QA.'),
    ('Type checking', 'Passed with no TypeScript errors.'),
    ('Focused tests', '17 tests passed across itinerary validation, pricing, and PDF export.'),
], [1.45, 5.0])

doc.add_heading('Remaining gaps from the broader audit', level=1)
paragraph('The fixes above address the requested editing and export behavior. The itinerary is still assembled from free-text schedule entries. It does not verify transfer time, opening hours, seasonal closures, or ticket availability. A selected package is copied into a customer PDF; the popup does not save a revised master package or a reusable customer draft.')
paragraph('The estimated total applies the same per person amount to adults and children. If child pricing differs, staff should edit the quote separately until a pricing breakdown is added. Trip dates and each day label are still independent text fields.')

doc.add_heading('Online comparison and next improvement', level=1)
paragraph('Wanderlog documents day-level route optimization. Google Routes can return travel distance and time, while Google Places exposes current and regular opening hours. The next product improvement remains a Trip Feasibility Check that validates each day before a customer PDF is sent. This requires structured place IDs, visit durations, travel legs, and a checked-at timestamp.')
paragraph('References: Wanderlog Optimize route — https://help.wanderlog.com/hc/en-us/articles/13545624787867-Optimize-route')
paragraph('Google Routes Compute Route Matrix — https://developers.google.com/maps/documentation/routes/compute-route-matrix-over')
paragraph('Google Places resource — https://developers.google.com/maps/documentation/places/web-service/reference/rest/v1/places')

doc.add_heading('Scope of verification', level=1)
paragraph('The popup was exercised against a temporary local preview with sample package data, then the preview route was removed. No live admin account, Firestore package, or customer record was changed. The exported PDF was checked with a focused automated test and a rendered sample cover. A signed-in production admin walkthrough remains useful before deployment.')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Hills Tourism itinerary builder reaudit  |  ')
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string(GRAY)
page = OxmlElement('w:fldSimple')
page.set(qn('w:instr'), 'PAGE')
footer._p.append(page)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = 'Hills Tourism Itinerary Builder Reaudit and Fix Report'
doc.core_properties.author = 'OpenAI Codex'
doc.save(OUTPUT)
print(OUTPUT.resolve())
